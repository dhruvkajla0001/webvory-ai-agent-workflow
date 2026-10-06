from pathlib import Path
from openpyxl import load_workbook
from .models import WorkflowSpec


def _split_steps(value: str) -> list[str]:
    return [x.strip() for x in value.split("→") if x.strip()]


def _split_tools(value: str) -> list[str]:
    return [x.strip() for x in value.split(";") if x.strip()]


def load_workflows(path: str | Path) -> dict[str, WorkflowSpec]:
    path = Path(path)
    wb = load_workbook(path, data_only=True)
    ws = wb["Workflows"]

    workflows: dict[str, WorkflowSpec] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        spec = WorkflowSpec(
            workflow_id=str(row[0]),
            workflow_name=str(row[1]),
            trigger=str(row[2]),
            inputs=str(row[3]),
            steps=_split_steps(str(row[4])),
            decision_logic=str(row[5]),
            tools_required=_split_tools(str(row[6])),
            expected_output=str(row[7]),
        )
        workflows[spec.workflow_id] = spec
    return workflows


def load_test_questions(path: str | Path) -> list[dict[str, str]]:
    path = Path(path)
    wb = load_workbook(path, data_only=True)
    ws = wb["Test_Questions"]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            rows.append({
                "workflow_id": str(row[0]),
                "test_request": str(row[1]),
                "what_to_check": str(row[2]),
            })
    return rows
